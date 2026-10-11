"""ATS parse-check engine: normalization, checks, scoring and round trip."""

import base64
import copy
import io
import json
import time
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.shared import Inches
from pydantic import ValidationError

from app.services.ats_parse import build_report, parse_check_bytes
from app.services.ats_parse.checks import find_sections, has_columns, run_checks, score
from app.services.ats_parse.extract import (
    MAX_COUNTED_PAGES,
    MAX_PARSE_CHECK_PAGES,
    ExtractedDocument,
    PageLayout,
    TextBox,
    extract_pdf,
)
from app.services.ats_parse.headings import ACADEMIC_SECTIONS, HEADINGS
from app.services.ats_parse.models import Check
from app.services.ats_parse.normalize import normalize, rejoin_letter_spacing, tokens
from app.services.ats_parse.roundtrip import (
    compute_roundtrip,
    expected_fields,
    kendall_fidelity,
)
from tests.ats_fixtures import fixture, john_doe, render
from tests.unit.test_page_fit import _pdf

pytestmark = pytest.mark.unit

MESSAGES = Path(__file__).resolve().parents[3] / "frontend" / "messages"


def _status(checks: list[Check], check_id: str) -> str:
    return next(c.status for c in checks if c.id == check_id)


def _text_doc(text: str) -> ExtractedDocument:
    return ExtractedDocument(kind="pdf", text=text, page_count=1)


def _docx(build: Any) -> bytes:
    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _form_xobject_pdf(text: str) -> bytes:
    """One page whose only text is drawn inside a form XObject (``/Fx Do``)."""
    form = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /XObject << /Fx 5 0 R >> >> /Contents 4 0 R >>",
        "<< /Length 9 >> stream\n/Fx Do\n\nendstream",
        "<< /Type /XObject /Subtype /Form /BBox [0 0 612 792] "
        "/Resources << /Font << /F1 6 0 R >> >> "
        f"/Length {len(form)} >> stream\n{form}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj {body} endobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


# -- normalize -----------------------------------------------------------------


def test_letter_spaced_headings_are_rejoined() -> None:
    assert rejoin_letter_spacing("S U M M A R Y") == "SUMMARY"
    assert normalize("E X P E R I E N C E\nAcme") == "experience\nacme"
    assert rejoin_letter_spacing("I am a dev") == "I am a dev"


def test_email_matching_is_linear_on_dotted_runs() -> None:
    # Unbounded, both patterns backtracked from every position: about a minute.
    text = "a." * 100_000
    start = time.perf_counter()
    assert tokens(text) == ["a"] * 100_000
    assert _status(run_checks(_text_doc(text), {}), "contact_info") == "fail"
    assert time.perf_counter() - start < 2
    assert tokens("Mail john.doe+cv@mail.example.com") == [
        "mail",
        "john.doe+cv@mail.example.com",
    ]


def test_supplementary_han_is_split_per_character() -> None:
    assert tokens("\U0002000b\U00020089") == ["\U0002000b", "\U00020089"]


def test_cjk_split_keeps_latin_runs_whole() -> None:
    assert tokens("Java开发 Python3工程师") == [
        "java",
        "开",
        "发",
        "python3",
        "工",
        "程",
        "师",
    ]
    data = {
        "personalInfo": {"name": "Li Wei"},
        "additional": {"technicalSkills": ["Java"]},
    }
    by_field = {
        f.field: f.status for f in compute_roundtrip(data, "Li Wei\nJava开发").fields
    }
    assert by_field["additional.technicalSkills"] == "found"


def test_normalize_casefolds_so_sharp_s_matches_ss() -> None:
    assert tokens("Straße") == tokens("STRASSE") == ["strasse"]


def test_normalize_strips_inline_html_and_keeps_real_hyphens() -> None:
    assert (
        normalize("<strong>Led</strong> open-\nsource work") == "led open-source work"
    )


# -- layout ----------------------------------------------------------------------


def _page(*boxes: tuple[float, float, float, float]) -> PageLayout:
    return PageLayout(600, 800, [TextBox(*b, "x") for b in boxes])


def test_columns_need_a_gutter_over_a_large_span() -> None:
    two_columns = _page(
        (20, 660, 360, 760),
        (20, 540, 360, 640),
        (20, 420, 360, 520),
        (400, 700, 580, 760),
        (400, 600, 580, 680),
        (400, 500, 580, 580),
    )
    dates_beside_titles = _page(
        (20, 700, 150, 740),
        (480, 700, 580, 740),
        (20, 650, 580, 690),
        (20, 600, 580, 640),
    )
    assert has_columns(two_columns)
    assert not has_columns(dates_beside_titles)


def test_right_aligned_dates_in_a_run_of_entries_are_not_a_column() -> None:
    # Entries without body text leave no full-width line to cross the gutter, so the
    # dates form a band of their own. They all end at one edge and are narrow.
    titles_and_dates = _page(
        *[(20, 700 - 60 * i, 330, 740 - 60 * i) for i in range(4)],
        *[(550 - 15 * (i % 3), 705 - 60 * i, 580, 725 - 60 * i) for i in range(4)],
    )
    assert not has_columns(titles_and_dates)


def test_a_ragged_sidebar_is_still_a_column_even_when_narrow() -> None:
    sidebar = _page(
        *[(20, 700 - 60 * i, 330, 740 - 60 * i) for i in range(4)],
        *[(420, 705 - 60 * i, 420 + 60 + 25 * i, 725 - 60 * i) for i in range(4)],
    )
    assert has_columns(sidebar)


def test_a_justified_main_column_beside_a_sidebar_is_still_a_column() -> None:
    # Boxes that end at one edge but are wide are text, not right-aligned metadata.
    justified = _page(
        *[(20, 700 - 60 * i, 150, 740 - 60 * i) for i in range(4)],
        *[(200, 700 - 60 * i, 580, 740 - 60 * i) for i in range(4)],
    )
    assert has_columns(justified)


@pytest.mark.parametrize(
    ("name", "template"),
    [("swiss-single-academic-3p", "swiss-single"), ("clean-academic-3p", "clean")],
)
async def test_single_column_academic_cv_over_three_pages_is_not_multi_column(
    name: str, template: str
) -> None:
    # Publications, grants and talks are entries with a title and a right-aligned
    # year (clean: "Pittsburgh, PA | Jul 2021 - Present") and no body text. Every
    # page of these single-column templates used to be flagged as columns.
    report = await parse_check_bytes(
        fixture(name), f"{name}.pdf", source="render", template=template
    )
    assert report.page_count == 3
    check = next(c for c in report.checks if c.id == "multi_column")
    assert check.status == "pass"
    assert check.params["pages"] == []
    assert check.params["expected_by_template"] is False


@pytest.mark.parametrize("name", ["swiss-single-academic-3p", "clean-academic-3p"])
async def test_academic_appointments_on_a_real_render_count_as_experience(
    name: str,
) -> None:
    report = await parse_check_bytes(fixture(name), f"{name}.pdf", source="render")
    check = next(c for c in report.checks if c.id == "section_headings")
    assert "experience" in check.params["found"]


async def test_cv_mode_lifts_the_page_count_warning_but_not_truncation() -> None:
    content = fixture("swiss-single-academic-3p")
    default = await parse_check_bytes(content, "a.pdf", source="render")
    cv = await parse_check_bytes(content, "a.pdf", source="render", cv_mode=True)
    assert _status(default.checks, "page_count") == "warn"
    assert _status(cv.checks, "page_count") == "pass"
    assert cv.page_count == 3
    truncated = ExtractedDocument(
        kind="pdf", text="x", page_count=MAX_PARSE_CHECK_PAGES + 1, truncated=True
    )
    checks = run_checks(truncated, {"cv_mode": True})
    assert _status(checks, "page_count") == "warn"


async def test_two_column_template_over_three_pages_is_still_detected() -> None:
    report = await parse_check_bytes(
        fixture("swiss-two-column-academic-3p"),
        "a.pdf",
        source="render",
        template="swiss-two-column",
    )
    assert report.page_count == 3
    check = next(c for c in report.checks if c.id == "multi_column")
    assert check.status == "fail"
    assert check.params["pages"] == [1]


@pytest.mark.parametrize(
    ("name", "status"),
    [
        ("single-column", "pass"),
        ("modern-single-column", "pass"),
        ("two-column", "fail"),
        ("modern-two-column", "fail"),
    ],
)
async def test_multi_column_on_real_renders(name: str, status: str) -> None:
    report = await parse_check_bytes(render(name), f"{name}.pdf")
    assert _status(report.checks, "multi_column") == status
    assert report.extractability == "full"


async def test_two_column_render_scrambles_reading_order() -> None:
    single = await parse_check_bytes(
        render("single-column"), "a.pdf", source="render", resume_data=john_doe()
    )
    double = await parse_check_bytes(
        render("two-column"), "b.pdf", source="render", resume_data=john_doe()
    )
    assert single.roundtrip.content_recall == 1.0
    assert single.roundtrip.order_fidelity == 1.0
    assert double.roundtrip.order_fidelity < 0.9
    assert double.overall_score < single.overall_score


async def test_modern_deedy_render_is_single_column_and_extracts_cleanly() -> None:
    # Real render of john_doe(): the roles and institutions are uppercase text, not
    # font small caps (which reach the text layer as private-use glyphs, as in clean),
    # the pipes are real text, and the read order is high.
    report = await parse_check_bytes(
        render("modern-deedy"),
        "d.pdf",
        source="render",
        resume_data=john_doe(),
        template="modern-deedy",
    )
    assert report.overall_score == 100
    assert _status(report.checks, "garbled_glyphs") == "pass"
    assert _status(report.checks, "multi_column") == "pass"
    assert report.roundtrip.content_recall == 1.0
    assert report.roundtrip.order_fidelity >= 0.99
    text = extract_pdf(render("modern-deedy")).text
    assert "STRIPE | SENIOR SOFTWARE ENGINEER" in text
    assert "Technical Skills: Go, Python, Kafka" in text


async def test_sparse_sidebar_still_counts_as_two_columns() -> None:
    # A short skills sidebar covers little of the page but still interleaves.
    report = await parse_check_bytes(
        fixture("swiss-two-column-sparse"),
        "a.pdf",
        source="render",
        resume_data=john_doe(),
        template="swiss-two-column",
    )
    assert _status(report.checks, "multi_column") == "fail"
    assert report.roundtrip.order_fidelity < 0.9


async def test_small_caps_glyphs_are_flagged_and_not_matched_elsewhere() -> None:
    # The clean template's small-caps titles encode "e" as U+F765, so
    # "Senior Software Engineer" in the header must not stand in for the job title.
    report = await parse_check_bytes(
        fixture("clean-small-caps"), "c.pdf", source="render", resume_data=john_doe()
    )
    assert _status(report.checks, "garbled_glyphs") == "warn"
    fields = {f.field: f.status for f in report.roundtrip.fields}
    assert fields["personalInfo.title"] == "found"
    assert fields["workExperience[0].title"] == "missing"
    assert _status(report.checks, "section_headings") == "pass"


async def test_vivid_single_render_is_fully_parseable_and_in_order() -> None:
    # A real Chromium render of john_doe() with the vivid-single template. Small caps
    # and an opacity tint would put private-use glyphs or an unread "image" in the
    # text layer; the single column must not be flagged as columns either.
    report = await parse_check_bytes(
        fixture("vivid-single"),
        "v.pdf",
        source="render",
        resume_data=john_doe(),
        template="vivid-single",
    )
    assert report.overall_score == 100
    for check_id in ("garbled_glyphs", "images", "multi_column", "section_headings"):
        assert _status(report.checks, check_id) == "pass", check_id
    multi = next(c for c in report.checks if c.id == "multi_column")
    assert multi.params["expected_by_template"] is False
    assert report.roundtrip.content_recall == 1.0
    assert report.roundtrip.order_fidelity == 1.0
    assert "John Doe" in report.extracted_text_preview


async def test_roundtrip_uses_the_template_contact_order() -> None:
    # The clean template renders location, phone, then email.
    report = await parse_check_bytes(
        fixture("clean-small-caps"),
        "c.pdf",
        source="render",
        resume_data=john_doe(),
        template="clean",
    )
    assert report.roundtrip.order_fidelity == 1.0
    names = [name for name, _ in expected_fields(john_doe(), "latex")[0]][:5]
    assert names[2:] == [
        "personalInfo.location",
        "personalInfo.phone",
        "personalInfo.email",
    ]


async def test_roundtrip_checks_header_links() -> None:
    # The render shows "linkedin.com/in/johndoe , github.com/johndoe".
    data = john_doe()
    data["personalInfo"]["linkedin"] = "www.linkedin.com/in/johndoe"
    data["personalInfo"]["github"] = "github.com/johndoe"
    text = extract_pdf(render("single-column")).text
    kept = compute_roundtrip(data, text, "swiss-single")
    statuses = {f.field: f.status for f in kept.fields}
    assert statuses["personalInfo.linkedin"] == "found"
    assert statuses["personalInfo.github"] == "found"
    assert kept.content_recall == 1.0 and kept.order_fidelity == 1.0

    dropped = compute_roundtrip(data, text.replace("linkedin.com/in/johndoe", ""))
    assert {f.field: f.status for f in dropped.fields}[
        "personalInfo.linkedin"
    ] == "missing"
    assert dropped.content_recall < 1.0


def test_roundtrip_links_follow_template_display_and_order() -> None:
    data = john_doe()
    data["personalInfo"] |= {
        "website": "https://www.jane.dev",  # a full URL is shown as typed
        "linkedin": "www.linkedin.com/in/jane",  # a bare host loses "www."
        "github": "github.com/jane",
    }
    expected = {
        None: ["email", "phone", "location", "website", "linkedin", "github"],
        "clean": ["location", "phone", "email", "linkedin", "github", "website"],
        "latex": ["location", "phone", "email", "linkedin", "github", "website"],
        "vivid": ["website", "linkedin", "github", "email", "phone", "location"],
        "vivid-single": ["website", "linkedin", "github", "email", "phone", "location"],
        "modern-deedy": ["email", "phone", "location", "website", "linkedin", "github"],
        "deedy": ["website", "linkedin", "github", "email", "phone", "location"],
    }
    for template, order in expected.items():
        fields = expected_fields(data, template)[0][2:8]
        assert [name.removeprefix("personalInfo.") for name, _ in fields] == order
    texts = dict(expected_fields(data)[0])
    assert texts["personalInfo.website"] == "https://www.jane.dev"
    assert texts["personalInfo.linkedin"] == "linkedin.com/in/jane"


def test_roundtrip_reads_entry_rows_left_column_first() -> None:
    # The render's "title | dates" over "company | location" rows read title,
    # company, dates, location.
    data = john_doe()
    for entry in data["workExperience"]:
        entry["location"] = "San Francisco, CA"
    text = extract_pdf(render("single-column")).text
    assert compute_roundtrip(data, text, "swiss-single").order_fidelity == 1.0


# Entry lines in the order pdfminer reads assets/pdf-templates/<template>.pdf.
TEMPLATE_ENTRY_TEXT = {
    "clean": "Acme Corp | Staff Engineer\nDublin | 2020 - 2023\n"
    "Trinity College | BSc Physics\n2016",
    "latex": "Acme Corp\nStaff Engineer\n2020 - 2023\nDublin\n"
    "Trinity College\nBSc Physics\n2016",
    "vivid": "Acme Corp | Staff Engineer\n2020 - 2023 | Dublin\n"
    "Trinity College\nBSc Physics\n2016",
    "vivid-single": "Acme Corp | STAFF ENGINEER\n2020 - 2023 | Dublin\n"
    "Trinity College | BSC PHYSICS\n2016",
    "modern-deedy": "Acme Corp | Staff Engineer\nDublin | 2020 - 2023\n"
    "BSc Physics\nTrinity College\n2016",
    "deedy": "ACME CORP | STAFF ENGINEER\n2020 - 2023 | Dublin\n"
    "TRINITY COLLEGE\nBSC PHYSICS\n2016",
    "swiss-two-column": "Staff Engineer\nAcme Corp • Dublin\n2020 - 2023\n"
    "Trinity College | 2016\nBSc Physics",
}


@pytest.mark.parametrize("template", TEMPLATE_ENTRY_TEXT)
def test_roundtrip_uses_the_template_entry_order(template: str) -> None:
    data = {
        "personalInfo": {"name": "Jane Roe"},
        "workExperience": [
            {
                "title": "Staff Engineer",
                "company": "Acme Corp",
                "location": "Dublin",
                "years": "2020 - 2023",
            }
        ],
        "education": [
            {"institution": "Trinity College", "degree": "BSc Physics", "years": "2016"}
        ],
    }
    text = "Jane Roe\n" + TEMPLATE_ENTRY_TEXT[template]
    assert compute_roundtrip(data, text, template).order_fidelity == 1.0
    assert compute_roundtrip(data, text).order_fidelity < 1.0


async def test_pdf_without_text_is_not_extractable() -> None:
    report = await parse_check_bytes(_pdf(1), "scan.pdf")
    assert report.extractability == "none"
    assert _status(report.checks, "text_layer") == "fail"


def test_pdf_page_count_stops_at_the_budget() -> None:
    doc = extract_pdf(_pdf(MAX_COUNTED_PAGES + 50))
    assert doc.page_count == MAX_COUNTED_PAGES
    assert doc.truncated
    assert len(doc.pages) == MAX_PARSE_CHECK_PAGES


async def test_pdf_text_inside_form_xobjects_is_extracted() -> None:
    line = "Jane Doe jane@example.com Senior Engineer at Acme since 2019"
    report = await parse_check_bytes(_form_xobject_pdf(line), "form.pdf")
    assert line in report.extracted_text_preview
    assert _status(report.checks, "text_layer") == "pass"


async def test_docx_tables_and_header_contact_are_flagged() -> None:
    def build(document: Any) -> None:
        document.sections[0].header.paragraphs[0].text = "jane@example.com 555-010-0100"
        document.add_paragraph("EXPERIENCE")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Acme Corp"
        table.cell(0, 1).text = "2020 - 2022"

    report = await parse_check_bytes(_docx(build), "resume.docx")
    assert _status(report.checks, "tables") == "warn"
    assert _status(report.checks, "header_footer_content") == "fail"
    assert _status(report.checks, "contact_info") == "fail"
    assert "Acme Corp | 2020 - 2022" in report.extracted_text_preview


async def test_docx_tabs_and_line_breaks_separate_words() -> None:
    def build(document: Any) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3))
        run = paragraph.add_run("Jane Doe")
        run.add_tab()
        run.add_text("Engineer")
        run.add_break()
        run.add_text("Acme")
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run("Go").add_tab()
        cell.paragraphs[0].add_run("Python")

    report = await parse_check_bytes(_docx(build), "resume.docx")
    lines = report.extracted_text_preview.split("\n")
    assert lines[:2] == ["Jane Doe\tEngineer", "Acme"]
    assert "Go\tPython" in lines


@pytest.mark.parametrize("part", ["first_page_header", "even_page_footer"])
async def test_docx_contact_in_first_or_even_page_margin_is_flagged(part: str) -> None:
    def build(document: Any) -> None:
        section = document.sections[0]
        section.different_first_page_header_footer = True
        document.settings.odd_and_even_pages_header_footer = True
        getattr(section, part).paragraphs[0].text = "jane@example.com"
        document.add_paragraph("EXPERIENCE")

    report = await parse_check_bytes(_docx(build), "resume.docx")
    check = next(c for c in report.checks if c.id == "header_footer_content")
    assert check.status == "fail"
    assert check.params["contact_only_in_margin"] is True


async def test_docx_header_image_is_counted() -> None:
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )

    def build(document: Any) -> None:
        header = document.sections[0].header
        header.paragraphs[0].add_run().add_picture(io.BytesIO(png))
        document.add_paragraph("Jane Doe jane@example.com")

    report = await parse_check_bytes(_docx(build), "resume.docx")
    check = next(c for c in report.checks if c.id == "images")
    assert check.status == "warn"
    assert check.params["count"] == 1


async def test_docx_table_cell_paragraphs_stay_separate_words() -> None:
    def build(document: Any) -> None:
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].paragraphs[0].text = "Senior"
        cells[0].add_paragraph("Engineer")
        cells[1].text = "2020"

    report = await parse_check_bytes(_docx(build), "resume.docx")
    assert "Senior\nEngineer | 2020" in report.extracted_text_preview
    assert "SeniorEngineer" not in report.extracted_text_preview


async def test_docx_contact_in_a_header_table_is_flagged() -> None:
    def build(document: Any) -> None:
        header = document.sections[0].header
        table = header.add_table(rows=1, cols=2, width=document.sections[0].page_width)
        table.cell(0, 0).text = "Jane Doe"
        table.cell(0, 1).text = "jane@example.com"
        document.add_paragraph("EXPERIENCE")

    report = await parse_check_bytes(_docx(build), "resume.docx")
    check = next(c for c in report.checks if c.id == "header_footer_content")
    assert check.status == "fail"
    assert check.params["contact_only_in_margin"] is True


# -- content ---------------------------------------------------------------------


def test_contact_info_found_late_is_a_warning() -> None:
    early = run_checks(
        _text_doc("Jane jane@x.com +1 415 555 0100\n" + "word " * 200), {}
    )
    late = run_checks(_text_doc("word " * 200 + "jane@x.com +1 415 555 0100"), {})
    assert _status(early, "contact_info") == "pass"
    assert _status(late, "contact_info") == "warn"


@pytest.mark.parametrize(
    ("text", "phone"),
    [
        ("Acme 2019 - 2022", False),
        ("Acme 03.2019 - 06.2022", False),
        ("2019 - 2022\n+44 20 7946 0958", True),
        ("(415) 555-0142", True),
        # Seven-digit local numbers.
        ("555-1234", True),
        ("5551234", True),
    ],
)
def test_contact_info_needs_a_phone_shaped_number(text: str, phone: bool) -> None:
    checks = run_checks(_text_doc(f"Jane jane@x.com\n{text}\n" + "word " * 50), {})
    check = next(c for c in checks if c.id == "contact_info")
    assert check.params["phone"] is phone
    assert check.status == ("pass" if phone else "warn")


@pytest.mark.parametrize(
    ("text", "locale", "status"),
    [
        ("Experiencia\nEducación\nHabilidades", "es", "pass"),
        ("Experience\nEducation\nSkills", "ja", "pass"),
        ("Experience\nSkills", None, "warn"),
        ("Experience with Python and teams across many offices", None, "fail"),
        # Short prose that starts with a synonym is not a heading.
        ("Experience in product engineering\nEducation\nSkills", None, "fail"),
        ("職歴\n学歴\n技術を活用したシステム開発", "ja", "warn"),
        ("職歴\n学歴\n技術スキル", "ja", "pass"),
        # Academic appointments stand in for experience.
        ("Academic Appointments\nEducation\nSkills", None, "pass"),
        ("Appointments\nEducation\nSkills", "en", "pass"),
        ("Research Experience\nEducation\nSkills", None, "pass"),
        ("Academic Appointments\nEducation", None, "warn"),
        ("Education\nSkills\nPublications", None, "fail"),
        # Prose that starts with an appointment synonym is not a heading.
        ("Appointments scheduling for clinics\nEducation\nSkills", None, "fail"),
    ],
)
def test_section_headings(text: str, locale: str | None, status: str) -> None:
    assert (
        _status(run_checks(_text_doc(text), {"locale": locale}), "section_headings")
        == status
    )


@pytest.mark.parametrize(
    ("line", "kind"),
    [
        ("Publications", "publications"),
        ("Selected Publications", "publications"),
        ("Peer-Reviewed Publications", "publications"),
        ("Peer-reviewed journal articles", "publications"),
        ("Journal Articles", "publications"),
        ("Conference Papers", "publications"),
        ("Peer-Reviewed Conference Papers", "publications"),
        ("Publications & Presentations", "publications"),
        ("Grants", "grants"),
        ("Grants & Funding", "grants"),
        ("Funding", "grants"),
        ("Grants and Funding", "grants"),
        ("Teaching", "teaching"),
        ("Invited Talks", "talks"),
        ("Presentations", "talks"),
        ("Service", "service"),
        ("Academic Service", "service"),
        ("Research", "research"),
    ],
)
def test_academic_headings_are_recognized(line: str, kind: str) -> None:
    found = find_sections(f"Jane Doe\n{line}\nSome text", None, ACADEMIC_SECTIONS)
    assert kind in found


@pytest.mark.parametrize(
    "line",
    [
        "Research interests include single-cell methods",
        "Teaching and mentoring graduate students across three labs",
        "Publications are listed on my website",
        "Service to the profession and the community at large",
        "Funding secured from several agencies",
        "Presentations at major venues worldwide this year",
    ],
)
def test_prose_starting_with_an_academic_word_is_not_a_heading(line: str) -> None:
    assert find_sections(line, None, ACADEMIC_SECTIONS) == []


def test_academic_sections_never_change_the_required_headings_report() -> None:
    # The report's found/missing lists name only the three required kinds; the UI maps
    # each to a translated label.
    doc = _text_doc("Academic Appointments\nEducation\nPublications\nGrants\nTeaching")
    check = next(c for c in run_checks(doc, {}) if c.id == "section_headings")
    assert check.params["found"] == ["experience", "education"]
    assert check.params["missing"] == ["skills"]


def test_headings_cover_every_frontend_locale_label() -> None:
    kinds = {"summary": "summary", "experience": "experience", "education": "education"}
    kinds |= {"projects": "projects", "skills": "skills"}
    for path in MESSAGES.glob("*.json"):
        locale = path.stem.split("-")[0]
        labels = json.loads(path.read_text(encoding="utf-8"))["resume"]["sections"]
        for label_key, kind in kinds.items():
            label = normalize(labels[label_key])
            assert any(
                label.startswith(normalize(s)) for s in HEADINGS[locale][kind]
            ), f"{path.name}: {label!r} not in HEADINGS[{locale!r}][{kind!r}]"


@pytest.mark.parametrize("skills_key", ["skills", "skillsOnly"])
def test_frontend_section_labels_read_as_headings(skills_key: str) -> None:
    # "Skills & Awards", "Habilidades y Premios", "技能与荣誉" and the like.
    for path in MESSAGES.glob("*.json"):
        labels = json.loads(path.read_text(encoding="utf-8"))["resume"]["sections"]
        text = "\n".join(labels[k] for k in ("experience", "education", skills_key))
        found = find_sections(text, path.stem)
        assert found == ["experience", "education", "skills"], (path.name, text)


def test_score_halves_warnings() -> None:
    checks = [
        Check(id="a", category="layout", severity="medium", status="fail"),
        Check(id="b", category="layout", severity="high", status="warn"),
        Check(id="c", category="content", severity="high", status="fail"),
    ]
    assert score(checks, {"layout"}) == 100 - 12 - 15
    assert score(checks, {"content"}) == 70


def test_multi_column_reports_whether_template_expects_it() -> None:
    doc = ExtractedDocument(kind="pdf", text="x", page_count=1, pages=[])
    report = build_report(doc, source="render", template="swiss-two-column")
    check = next(c for c in report.checks if c.id == "multi_column")
    assert check.params["expected_by_template"] is True


def test_deedy_columns_are_expected_by_the_template() -> None:
    doc = ExtractedDocument(kind="pdf", text="x", page_count=1, pages=[])
    report = build_report(doc, source="render", template="deedy")
    check = next(c for c in report.checks if c.id == "multi_column")
    assert check.params["expected_by_template"] is True


# -- round trip --------------------------------------------------------------------


def _ordered_text(data: dict[str, Any]) -> list[str]:
    return [text for _, text in expected_fields(data)[0]]


def test_kendall_fidelity_bounds() -> None:
    assert kendall_fidelity([1, 2, 3, 4]) == 1.0
    assert kendall_fidelity([4, 3, 2, 1]) == 0.0


def test_roundtrip_in_order_and_shuffled() -> None:
    data = john_doe()
    parts = _ordered_text(data)
    exact = compute_roundtrip(data, "\n".join(parts))
    reverse = compute_roundtrip(data, "\n".join(reversed(parts)))
    assert exact.content_recall == 1.0 and exact.order_fidelity == 1.0
    assert reverse.content_recall == 1.0 and reverse.order_fidelity < 0.2


def test_roundtrip_reports_missing_and_hidden() -> None:
    data = john_doe()
    dropped = data["workExperience"][0]["description"][1]
    text = "\n".join(p for p in _ordered_text(data) if p != dropped)
    hidden = copy.deepcopy(data)
    for meta in hidden["sectionMeta"]:
        if meta["key"] == "education":
            meta["isVisible"] = False

    result = compute_roundtrip(data, text)
    by_field = {f.field: f.status for f in result.fields}
    assert by_field["workExperience[0].description[1]"] == "missing"
    assert result.content_recall < 1.0

    hidden_result = compute_roundtrip(hidden, text)
    statuses = {f.field: f.status for f in hidden_result.fields}
    assert statuses["section.education"] == "hidden"
    assert not any(name.startswith("education[") for name in statuses)


def test_roundtrip_word_satisfies_one_field_only() -> None:
    data = {
        "personalInfo": {"name": "Jane Roe", "title": "Python Developer"},
        "additional": {"technicalSkills": ["Python"]},
    }
    skill = "additional.technicalSkills"
    omitted = compute_roundtrip(data, "Jane Roe\nPython Developer")
    listed = compute_roundtrip(data, "Jane Roe\nPython Developer\nSkills\nPython")
    assert {f.field: f.status for f in omitted.fields}[skill] == "missing"
    assert {f.field: f.status for f in listed.fields}[skill] == "found"
    assert listed.content_recall == 1.0


def test_roundtrip_checks_rendered_locations() -> None:
    data = john_doe()
    data["workExperience"][0]["location"] = "Dublin, Ireland"
    data["customSections"]["volunteer"] = {
        "sectionType": "itemList",
        "items": [{"title": "Mentor", "subtitle": "Code Club", "location": "Leeds"}],
    }
    data["sectionMeta"].append(
        {
            "id": "volunteer",
            "key": "volunteer",
            "displayName": "Volunteering",
            "sectionType": "itemList",
            "isDefault": False,
            "order": 9,
        }
    )
    text = "\n".join(_ordered_text(data))
    assert "Dublin, Ireland" in text and "Leeds" in text
    dropped = text.replace("Dublin, Ireland", "").replace("Leeds", "")
    by_field = {f.field: f.status for f in compute_roundtrip(data, dropped).fields}
    assert by_field["workExperience[0].location"] == "missing"
    assert by_field["customSections.volunteer.items[0].location"] == "missing"


def test_roundtrip_validates_stored_data() -> None:
    with pytest.raises(ValidationError):
        compute_roundtrip({"workExperience": ["bogus"]}, "text")
    # Data without sectionMeta renders the default sections, as the print page does.
    legacy = {"workExperience": [{"title": "Staff Engineer", "company": "Acme"}]}
    fields = {f.field for f in compute_roundtrip(legacy, "Staff Engineer Acme").fields}
    assert {"workExperience[0].title", "workExperience[0].company"} <= fields
