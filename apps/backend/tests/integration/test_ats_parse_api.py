"""ATS parse-check endpoints: uploads, rendered resumes and error mapping."""

import io
from collections.abc import AsyncIterator
from typing import Any
from unittest.mock import AsyncMock
from urllib.parse import parse_qs, urlsplit

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient

from app import services
from app.database import Database
from app.main import app
from app.pdf import PDFRenderError
from app.routers import resumes as resumes_router
from app.routers.resumes import MAX_FILE_SIZE
from tests.ats_fixtures import john_doe, render

pytestmark = pytest.mark.integration

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
async def client() -> AsyncIterator[AsyncClient]:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        yield c


def _statuses(report: dict[str, Any]) -> dict[str, str]:
    return {check["id"]: check["status"] for check in report["checks"]}


async def _resume(db: Database, data: dict[str, Any] | None = None) -> str:
    resume = await db.create_resume(
        content="x", processed_data=data, processing_status="ready"
    )
    return resume["resume_id"]


async def test_upload_pdf_reports_checks(client: AsyncClient) -> None:
    res = await client.post(
        "/api/v1/ats/parse-check",
        files={"file": ("cv.pdf", render("two-column"), PDF)},
    )
    assert res.status_code == 200
    body = res.json()
    assert body["source"] == "upload" and body["roundtrip"] is None
    assert _statuses(body)["multi_column"] == "fail"
    assert "JOHN DOE" in body["extracted_text_preview"]


async def test_upload_docx(client: AsyncClient) -> None:
    document = Document()
    document.add_paragraph("Jane Doe jane@example.com +1 415 555 0100")
    document.add_paragraph("Experience")
    buffer = io.BytesIO()
    document.save(buffer)
    res = await client.post(
        "/api/v1/ats/parse-check",
        files={"file": ("cv.docx", buffer.getvalue(), DOCX)},
    )
    assert res.status_code == 200
    assert _statuses(res.json())["tables"] == "pass"


@pytest.mark.parametrize(
    ("name", "content", "mime", "status"),
    [
        ("cv.doc", b"\xd0\xcf\x11\xe0", "application/msword", 400),
        ("cv.txt", b"hello", "text/plain", 400),
        ("cv.pdf", b"", PDF, 400),
        ("cv.pdf", b"x" * (MAX_FILE_SIZE + 1), PDF, 413),
        ("cv.pdf", b"%PDF-1.4 broken", PDF, 422),
    ],
)
async def test_upload_rejections(
    client: AsyncClient, name: str, content: bytes, mime: str, status: int
) -> None:
    res = await client.post(
        "/api/v1/ats/parse-check", files={"file": (name, content, mime)}
    )
    assert res.status_code == status


async def test_rendered_resume_matches_pdf_route_and_round_trips(
    client: AsyncClient, isolated_db: Database, monkeypatch: pytest.MonkeyPatch
) -> None:
    resume_id = await _resume(isolated_db, john_doe())
    check_render = AsyncMock(return_value=render("two-column"))
    route_render = AsyncMock(return_value=render("two-column"))
    monkeypatch.setattr(services.ats_parse, "render_resume_pdf", check_render)
    monkeypatch.setattr(resumes_router, "render_resume_pdf", route_render)
    settings = {"template": "swiss-two-column", "pageSize": "LETTER", "lang": "en"}

    res = await client.post(f"/api/v1/resumes/{resume_id}/parse-check", json=settings)
    assert res.status_code == 200
    body = res.json()
    assert body["source"] == "render" and body["template"] == "swiss-two-column"
    multi = next(c for c in body["checks"] if c["id"] == "multi_column")
    assert multi["status"] == "fail" and multi["params"]["expected_by_template"]
    assert body["roundtrip"]["content_recall"] > 0.9
    assert body["roundtrip"]["order_fidelity"] < 0.9

    await client.get(f"/api/v1/resumes/{resume_id}/pdf", params=settings)
    check_url = urlsplit(check_render.await_args.args[0])
    route_url = urlsplit(route_render.await_args.args[0])
    assert check_url.path == route_url.path
    assert parse_qs(check_url.query) == parse_qs(route_url.query)
    assert check_render.await_args.kwargs == route_render.await_args.kwargs


async def test_rendered_resume_errors(
    client: AsyncClient, isolated_db: Database, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert (await client.post("/api/v1/resumes/nope/parse-check")).status_code == 404
    unprocessed = await _resume(isolated_db)
    res = await client.post(f"/api/v1/resumes/{unprocessed}/parse-check")
    assert res.status_code == 422

    ready = await _resume(isolated_db, john_doe())
    monkeypatch.setattr(
        services.ats_parse,
        "render_resume_pdf",
        AsyncMock(side_effect=PDFRenderError("chromium detail")),
    )
    res = await client.post(f"/api/v1/resumes/{ready}/parse-check")
    assert res.status_code == 503
    assert "chromium detail" not in res.text
