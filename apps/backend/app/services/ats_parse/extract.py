"""ATS-style text extraction with layout data, run under the parser's bounded worker."""

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.oxml.ns import qn
from pdfminer.high_level import extract_pages
from pdfminer.layout import LAParams, LTContainer, LTFigure, LTImage, LTTextContainer
from pdfminer.pdfpage import PDFPage

from app.services.parser import (
    MAX_PDF_PREFLIGHT_PAGES,
    DocumentValidationError,
    _validate_docx_container,
    _validate_extracted_text,
    _validate_pdf_container,
    run_bounded_document_job,
)

MAX_PARSE_CHECK_PAGES = 10
# Page counting stops here, so a huge page tree costs a bounded walk. The
# container preflight shares the budget and rejects larger PDFs first, so this
# cap only matters when extract_pdf is called without the preflight.
MAX_COUNTED_PAGES = MAX_PDF_PREFLIGHT_PAGES
SUPPORTED_SUFFIXES = (".pdf", ".docx")

# pdfminer warns per glyph on Chromium fonts ("Could not get FontBBox").
logging.getLogger("pdfminer").setLevel(logging.ERROR)


@dataclass
class TextBox:
    x0: float
    y0: float
    x1: float
    y1: float
    text: str


@dataclass
class PageLayout:
    width: float
    height: float
    boxes: list[TextBox]
    images: int = 0

    @property
    def text(self) -> str:
        return "\n".join(box.text for box in self.boxes)


@dataclass
class ExtractedDocument:
    kind: Literal["pdf", "docx"]
    text: str
    page_count: int | None = None
    truncated: bool = False
    pages: list[PageLayout] = field(default_factory=list)
    images: int = 0
    tables: int = 0
    columns: int = 1
    header_footer_text: str = ""
    textbox_text: str = ""


def _count_images(item: LTContainer) -> int:
    count = 0
    for child in item:
        if isinstance(child, LTImage):
            count += 1
        elif isinstance(child, LTFigure):
            count += _count_images(child) or 1
    return count


def _text_boxes(item: LTContainer) -> list[TextBox]:
    """Text containers in layout order, including those inside form XObjects."""
    boxes: list[TextBox] = []
    for child in item:
        if isinstance(child, LTTextContainer):
            text = child.get_text().strip()
            if text:
                boxes.append(TextBox(child.x0, child.y0, child.x1, child.y1, text))
        elif isinstance(child, LTFigure):
            boxes.extend(_text_boxes(child))
    return boxes


def extract_pdf(content: bytes) -> ExtractedDocument:
    """pdfminer default layout order, the same reading MarkItDown and simple ATSs use.

    ``all_texts`` lays out text drawn inside form XObjects too; without it that
    text stays as loose characters and never reaches the report.
    """
    page_count = sum(
        1 for _ in PDFPage.get_pages(io.BytesIO(content), maxpages=MAX_COUNTED_PAGES)
    )
    pages: list[PageLayout] = []
    for page in extract_pages(
        io.BytesIO(content),
        laparams=LAParams(all_texts=True),
        maxpages=MAX_PARSE_CHECK_PAGES,
    ):
        pages.append(
            PageLayout(page.width, page.height, _text_boxes(page), _count_images(page))
        )
    text = "\n\n".join(page.text for page in pages)
    return ExtractedDocument(
        kind="pdf",
        text=text,
        page_count=page_count,
        truncated=page_count > MAX_PARSE_CHECK_PAGES,
        pages=pages,
        images=sum(page.images for page in pages),
    )


# Run content in document order. Only run-level tabs count: w:pPr/w:tabs/w:tab
# defines tab stops and is not text.
_RUN_CONTENT = "w:r/*[self::w:t or self::w:tab or self::w:br or self::w:cr]"
_RUN_BREAKS = {qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n"}


def _run_text(element: Any, path: str) -> str:
    """Run text with tabs and line breaks kept, so "Name<tab>Title" stays two words."""
    return "".join(
        _RUN_BREAKS.get(node.tag, node.text or "") for node in element.xpath(path)
    )


def _xml_text(element: Any) -> str:
    """Run text with one line per paragraph, so cell paragraphs never run together."""
    paragraphs: dict[Any, list[str]] = {}
    for node in element.xpath(f".//{_RUN_CONTENT}"):
        paragraph = next(node.iterancestors(qn("w:p")), None)
        paragraphs.setdefault(paragraph, []).append(
            _RUN_BREAKS.get(node.tag, node.text or "")
        )
    return "\n".join("".join(parts) for parts in paragraphs.values()).strip()


def _table_lines(table: Any) -> list[str]:
    """One line per row, its non-empty cells joined by " | "."""
    lines: list[str] = []
    for row in table.xpath("./w:tr"):
        cells = [_xml_text(cell) for cell in row.xpath("./w:tc")]
        lines.append(" | ".join(c for c in cells if c))
    return lines


def _count_docx_images(element: Any) -> int:
    return len(element.xpath(".//w:drawing")) + len(element.xpath(".//w:pict"))


def extract_docx(content: bytes) -> ExtractedDocument:
    """Body paragraphs and tables in document order; headers, footers and text boxes kept apart."""
    document = Document(io.BytesIO(content))
    body = document.element.body
    lines: list[str] = []
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(
                _run_text(child, f".//{_RUN_CONTENT}[not(ancestor::w:txbxContent)]")
            )
        elif tag == "tbl":
            lines.extend(_table_lines(child))
    header_footer: list[str] = []
    columns = 1
    images = _count_docx_images(body)
    for section in document.sections:
        # Word keeps first-page and even-page variants in parts of their own.
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if not part.is_linked_to_previous:
                header_footer.extend(p.text for p in part.paragraphs if p.text.strip())
                for table in part.tables:
                    header_footer.extend(t for t in _table_lines(table._tbl) if t)
                images += _count_docx_images(part._element)
        for num in section._sectPr.xpath("./w:cols/@w:num"):
            columns = max(columns, int(num))
    textboxes = [_xml_text(box) for box in body.xpath(".//w:txbxContent")]
    return ExtractedDocument(
        kind="docx",
        text="\n".join(line for line in lines if line.strip()),
        images=images,
        tables=len(document.tables),
        columns=columns,
        header_footer_text="\n".join(header_footer),
        textbox_text="\n".join(t for t in textboxes if t),
    )


def _extract_sync(content: bytes, filename: str) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentValidationError("Parse check supports PDF and DOCX files.")
    # Validate in memory: the resume never touches disk, so nothing is left behind.
    if suffix == ".pdf":
        _validate_pdf_container(io.BytesIO(content))
        result = extract_pdf(content)
    else:
        _validate_docx_container(io.BytesIO(content))
        result = extract_docx(content)
    _validate_extracted_text(result.text)
    return result


async def extract_document(content: bytes, filename: str) -> ExtractedDocument:
    return await run_bounded_document_job(_extract_sync, content, filename)
